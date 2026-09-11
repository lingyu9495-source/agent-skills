#!/usr/bin/env python3
"""docx 门面重建：正式封面 + 签署页 + 自动目录域
用法: python3 face_docx.py 输入.docx 输出.docx "项目全称" "编制单位" "文件编号"
"""
import sys, re, copy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, DST = sys.argv[1], sys.argv[2]
PNAME = sys.argv[3] if len(sys.argv) > 3 else "××××项目"
ORG   = sys.argv[4] if len(sys.argv) > 4 else "××××"
CODE  = sys.argv[5] if len(sys.argv) > 5 else "KY-2026-×××"

src_doc = Document(SRC)
out = Document()

# 页面
sec = out.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.6)

def set_font(run, cn="仿宋_GB2312", en="Times New Roman", size=12, bold=False, color=None):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    r = run._element.get_or_add_rPr().get_or_add_rFonts()
    r.set(qn("w:eastAsia"), cn)
    if color: run.font.color.rgb = RGBColor(*color)

def para(text="", cn="仿宋_GB2312", size=12, bold=False, align=None, before=0, after=6, line=None):
    p = out.add_paragraph()
    if align is not None: p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line: pf.line_spacing = Pt(line)
    if text:
        r = p.add_run(text); set_font(r, cn, size=size, bold=bold)
    return p

def page_break():
    out.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def field(par, instr):
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    par._p.append(b); par._p.append(i); par._p.append(e)

def style_heading(name, size, cn, before=18, after=10):
    st = out.styles[name]
    st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), cn)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

style_heading("Heading 1", 16, "黑体", 20, 12)
style_heading("Heading 2", 14, "黑体")
style_heading("Heading 3", 13, "楷体_GB2312")
style_heading("Heading 4", 12, "仿宋_GB2312")
nm = out.styles["Normal"]
nm.font.name = "Times New Roman"; nm.font.size = Pt(12)
nm.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "仿宋_GB2312")
nm.paragraph_format.line_spacing = Pt(26)

# ---------- 封面 ----------
for _ in range(5): para("", after=0)
para(PNAME.split("可行性研究")[0], cn="黑体", size=26, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, line=40)
para("可 行 性 研 究 报 告", cn="黑体", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, line=36)
for _ in range(9): para("", after=0)
para(f"编制单位：{ORG}", cn="仿宋_GB2312", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
para("二〇二六年九月", cn="仿宋_GB2312", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
page_break()

# ---------- 签署页 ----------
for _ in range(3): para("", after=0)
para(f"文件编号：{CODE}", size=12, after=20)
para("资料密级：内部资料 · 注意保存", size=12, after=40)
tbl = out.add_table(rows=6, cols=2)
tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
signs = [("编制人", ""), ("审核人", ""), ("审定人", ""), ("项目负责人", ""), ("编制单位（盖章）", ""), ("日    期", "")]
for i, (k, v) in enumerate(signs):
    c0, c1 = tbl.rows[i].cells
    c0.width, c1.width = Cm(4.5), Cm(9)
    p0 = c0.paragraphs[0]; r0 = p0.add_run(k); set_font(r0, "黑体", size=12)
page_break()

# ---------- 目录 ----------
para("目  录", cn="黑体", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
fld = out.add_paragraph(); field(fld, 'TOC \\o "1-3" \\h \\z \\u')
page_break()

# ---------- 正文搬运：跳过源文档的平铺封面文字，从正文起全部搬 ----------
body_src = src_doc.element.body
skip_until_vol = False
copied = 0
for child in list(body_src):
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        # 取该段文字判断是否正文起点
        texts = child.findall('.//' + qn('w:t'))
        txt = ''.join(t.text or '' for t in texts).strip()
        if txt.startswith('第一卷') or txt.startswith('第1卷') or '卷 总报告' in txt:
            skip_until_vol = True  # 从这一段开始保留
        if not skip_until_vol:
            continue  # 跳过封面区旧段落
        out.element.body.append(child)
        copied += 1
    elif tag == 'tbl':
        if skip_until_vol:
            out.element.body.append(child)
            copied += 1
    elif tag == 'sectPr':
        pass  # 用输出文档自己的节属性

# 页眉页脚
hdr = out.sections[0].header
hp = hdr.paragraphs[0]; hp.text = ""
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = hp.add_run(f"{PNAME.split('可行性研究')[0]} 可行性研究报告")
set_font(r, "黑体", size=9)
ftr = out.sections[0].footer
fp = ftr.paragraphs[0]; fp.text = ""
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r1 = fp.add_run("第 "); set_font(r1, size=9)
rp = fp.add_run(); field(fp, "PAGE"); set_font(rp, size=9)
r2 = fp.add_run(" 页 共 "); set_font(r2, size=9)
rn = fp.add_run(); field(fp, "NUMPAGES"); set_font(rn, size=9)
r3 = fp.add_run(" 页"); set_font(r3, size=9)

out.save(DST)
print(f"OK -> {DST} 搬运正文元素 {copied}")
