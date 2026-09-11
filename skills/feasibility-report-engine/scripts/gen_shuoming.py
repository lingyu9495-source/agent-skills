#!/usr/bin/env python3
"""《编制说明与资料需求清单》生成器 —— 可研交付三件套之一（1-2页简短说明）

用前必装: pip install python-docx
用法:
  python3 gen_shuoming.py 输出.docx "项目名称" ["总投资X亿元" "配套文件名1、文件名2"]

作用：客户拿到报告后，用这份1-2页的说明看懂：
  ①本阶段交付什么、什么深度（前期论证版≠报批版，管理预期）
  ②测算套表怎么用（蓝=输入）
  ③要升级正式版需补哪些资料（分批清单 = 二次收费抓手）
  ④编制口径与依据（发改委2023大纲+经济评价第三版，立专业人设）

立场铁律：编制单位/日期落款留空（客户自己署名）；无任何 AI 痕迹。
内容模板基于真实海洋产业大单实战，可自由裁剪；列表内容按项目定制。
"""
import sys
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = sys.argv[1] if len(sys.argv) > 1 else "编制说明与资料需求清单.docx"
PNAME = sys.argv[2] if len(sys.argv) > 2 else "××××项目"
INV = sys.argv[3] if len(sys.argv) > 3 else "约 ×× 亿元"
FILES = sys.argv[4] if len(sys.argv) > 4 else f"《{PNAME}可行性研究报告》（前期论证版）、《项目投资测算套表》"

out = Document()
sec = out.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.left_margin = Cm(3.0); sec.right_margin = Cm(2.6)

def sf(run, cn="仿宋_GB2312", size=12, bold=False):
    run.font.name = "Times New Roman"; run.font.size = Pt(size); run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)

def para(text="", cn="仿宋_GB2312", size=12, bold=False, align=None, after=6, before=0):
    p = out.add_paragraph()
    if align: p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = Pt(24)
    if text:
        r = p.add_run(text); sf(r, cn, size, bold)
    return p

def heading(t, lvl=1):
    p = out.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if lvl==1 else 10); p.paragraph_format.space_after = Pt(6)
    sizes = {1:(16,"黑体"),2:(14,"黑体"),3:(13,"楷体_GB2312")}
    sz, cn = sizes.get(lvl, (12,"黑体"))
    r = p.add_run(t); sf(r, cn, sz, True)

# ===== 正文 =====
para("编制说明与资料需求清单", cn="黑体", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
para(f"项目名称：{PNAME}", after=4)
para(f"配套文件：{FILES}", after=16)

heading("一、本阶段成果说明", 2)
heading("1. 成果定位", 3)
para(f"本项目为投资规模{INV}的产业项目，投资决策所需的部分基础条件（股权安排、审批手续、市场订单、融资结构等）尚在推进过程中。鉴于此，本阶段先行交付《可行性研究报告（前期论证版）》，重点完成以下工作：")
para("1. 战略与政策论证：论证项目与国家战略方向、区域产业政策的契合性；", after=3)
para("2. 产业与市场框架论证：梳理各业务板块的市场空间、竞争格局与协同逻辑；", after=3)
para("3. 建设与运营方案框架：明确总体布局、分期思路与商业模式；", after=3)
para("4. 财务评价框架与模拟测算：在现有参数基础上建立测算模型，给出基准情景下的财务指标区间，为后续精确测算搭好框架。", after=6)
para("本阶段成果的定位是把方向论证清楚、把测算框架搭好，供项目前期方向研判与资料准备使用，不等同于可用于立项报批或银行审贷的最终可研报告。", after=10)

heading("2. 测算套表说明", 3)
para("随附《项目投资测算套表》已将投资估算、资金筹措、分年度收入成本、现金流量与核心财务指标（财务内部收益率、财务净现值、投资回收期）的计算公式全部搭建完毕：")
para("1. 蓝色单元格为输入项（集中在参数主控工作表），填入项目实际数据后全表自动联动计算；", after=3)
para("2. 当前表中预置参数为按行业基准设定的示例值，用于演示测算逻辑，正式结论须以实际尽调数据回填；", after=3)
para("3. 套表与报告财务评价章节口径一致，可互相印证。", after=10)

heading("二、后续深化需补充的资料清单", 2)
para("为使报告达到立项报批、银行审贷所需深度，建议按以下三批补充资料，以便在此基础上完成正式版编制（含十七张标准测算报表）。")
heading("第一批：量价与投资基础数据（直接影响财务测算）", 3)
para("1. 各业务板块投资分项明细：工程费用、设备购置、安装工程、其他费用、预备费、建设期利息、流动资金；", after=3)
para("2. 各板块达产收入模型所需量价参数：产能、售价、出租率/周转率、成本结构；", after=3)
para("3. 建设期与投产爬坡时间表；", after=3)
para("4. 拟采用的融资结构：资本金比例、意向融资渠道及额度、融资成本假设。", after=8)
heading("第二批：主体与合规确权文件（影响项目能否落地）", 3)
para("1. 投资主体对相关目标公司/资产的收购、增资安排进展，交易对价与时间表；", after=3)
para("2. 现有土地、海域、岸线等权属证照与规划性质清单；", after=3)
para("3. 与地方政府相关协议的最新状态或接续安排；", after=3)
para("4. 涉行业审批事项办理状态（用地、环评、能评、稳评、岸线/海域、安全审查等）；如尚未启动，可协助梳理审批路径与时间预估。", after=8)
heading("第三批：市场与订单实证材料（增强报告说服力）", 3)
para("1. 各核心板块已接触客户的合作意向、框架协议或需求函；", after=3)
para("2. 区域内同类项目与竞争对手情况；", after=3)
para("3. 产业链上下游节点与外部合作方的投资边界与合作方式说明。", after=10)

heading("三、下阶段工作安排", 2)
para("第一、二批资料齐备后，可在此基础上完成正式版可研报告及十七张标准测算报表，预计编制周期为资料齐备后十五至二十个工作日；第三批市场实证材料到位后，可进一步形成专项尽调与决策支持材料。", after=10)

heading("四、编制口径说明", 2)
para("本报告编制依据为国家发展改革委《投资项目可行性研究报告编写大纲及说明》（发改投资规〔2023〕304号）及《建设项目经济评价方法与参数（第三版）》。报告引用数据均注明来源；测算部分在基础数据尚未完全落实处，采用行业基准与审慎假设，相关假设已在正文相应位置注明，供复核指正。", after=12)
para("如对报告内容或测算套表使用有任何疑问，欢迎随时沟通。", after=30)

# 落款留空（由客户自行署名）
para("", after=0); para("", after=0)
para("编制单位：", after=6)
para("年    月    日", after=0)

cp = out.core_properties
cp.author = ""; cp.last_modified_by = ""; cp.title = "编制说明与资料需求清单"; cp.comments = ""
out.save(OUT)
print(f"✓ 编制说明已生成 → {OUT}")
